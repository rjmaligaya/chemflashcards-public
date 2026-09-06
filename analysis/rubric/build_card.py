#!/usr/bin/env python3
"""Build the simplified 7-row Q-full coding card as a fillable workbook."""
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

ARIAL = "Arial"
H1 = Font(name=ARIAL, size=14, bold=True)
H2 = Font(name=ARIAL, size=11, bold=True)
BODY = Font(name=ARIAL, size=10)
BODYB = Font(name=ARIAL, size=10, bold=True)
MONO = Font(name=ARIAL, size=10)
BLUE = Font(name=ARIAL, size=10, color="0000FF")
GREY = Font(name=ARIAL, size=9, color="666666")
HDRFILL = PatternFill("solid", fgColor="D9D9D9")
FEATFILL = PatternFill("solid", fgColor="DCE6F1")
FORMFILL = PatternFill("solid", fgColor="FDE9D9")
INFILL = PatternFill("solid", fgColor="FFFF00")
EXFILL = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

# ── source data ─────────────────────────────────────────────────────────────
s = pd.read_csv("qfull_scored.csv")


# Collapses the suffix and substituent locant scores into one LOCANTS mark.
def locants(r):
    v = [r.suffix_locant, r.subst_locant]
    v = [x for x in v if pd.notna(x)]
    return "NA" if not v else int(min(v))


# Renders one slot score as an integer, or "NA" when it does not apply.
def cell(v):
    return "NA" if pd.isna(v) else int(v)


s = s.sort_values(["full_name", "pid"]).reset_index(drop=True)
rows = []
for _, r in s.iterrows():
    rows.append(dict(
        pid=r.pid, schedule=r.schedule, list=r.group_id, item_id=r.item_id,
        target=r.full_name,
        response=("(blank)" if pd.isna(r.raw_answer) else str(r.raw_answer)),
        PARENT=cell(r.root), CLASS=cell(r.suffix), SUBST=cell(r.subst_id),
        LOCANTS=locants(r), ORDER=cell(r.assembly), SPELLING=cell(r.ortho),
        SEPARATORS=cell(r.sep), tier=int(r.tier),
        tags=("" if pd.isna(r.tags) else r.tags),
        timeout=int(r.timed_out), adjudicate=("Y" if int(r.adjudicate) else "N"),
        note=("" if pd.isna(r.coder_note) else r.coder_note)))
coded = pd.DataFrame(rows)

SLOTS = ["PARENT", "CLASS", "SUBST", "LOCANTS", "ORDER", "SPELLING", "SEPARATORS"]

wb = Workbook()

# ════════════════════════════════════════════════════════════════════════════
# 1. HOW TO USE
# ════════════════════════════════════════════════════════════════════════════
ws = wb.active
ws.title = "How to use"
ws.sheet_view.showGridLines = False


# Writes one worksheet cell with the given font, wrapping, fill, and border.
def put(ws, r, c, v, font=BODY, wrap=False, fill=None, border=False):
    x = ws.cell(row=r, column=c, value=v)
    x.font = font
    x.alignment = Alignment(wrap_text=wrap, vertical="top")
    if fill:
        x.fill = fill
    if border:
        x.border = BOX
    return x


put(ws, 1, 1, "Q-full coding card: how to use", H1)
put(ws, 2, 1, "Segment-level scoring of the Study 2 delayed-test whole-name items. "
              "Slot map v1, simplified to seven rows. Draft for lab discussion, not yet pre-specified.", GREY)

put(ws, 4, 1, "The seven rows", H2)
put(ws, 5, 1, "For every response you ask the same seven questions. You do not need a per-item slot map: "
              "the target name itself tells you which rows apply. Mark 1 for hit, 0 for miss, NA where the "
              "target does not call for that row.", BODY, wrap=True)
ws.merge_cells("A5:G5")
ws.row_dimensions[5].height = 30

hdr = ["Row", "Question", "Mark NA when", "1 (hit) looks like", "0 (miss) looks like"]
for j, h in enumerate(hdr, start=1):
    put(ws, 7, j, h, H2, fill=HDRFILL, border=True)

card_rows = [
    ("PARENT", "Is the parent right? Chain length, and ring identity where the target is cyclic or aromatic.",
     "never, always applicable", "hept in 1-chloroheptane", "2-bromo-propane for 2-bromobutane"),
    ("CLASS", "Is the functional class right? The principal suffix, or the class word in a functional-class name.",
     "never, always applicable", "-amide in pentanamide", "pentaneol for pentanamide"),
    ("SUBST", "Is the substituent right? The prefix vocabulary, or the branch descriptor and alkyl group.",
     "the target carries no substituent", "fluoro in 1-fluorohexane", "bromotertbutane for iso-butyl bromide"),
    ("LOCANTS", "Are all required locants present and correct in value?",
     "the target requires no locant", "the 2 in 2-bromobutane", "fluoro-hexane, locant missing"),
    ("ORDER", "Are the segments in the right sequence, with the right word structure? Includes the ester two-word rule and the -an- linker.",
     "nothing was retrieved (tier 0)", "methyl octanoate", "octyl methanoate, hept-2-one"),
    ("SPELLING", "Is every retrieved segment spelled correctly?",
     "nothing was retrieved (tier 0)", "fluoro", "1-flourohexane"),
    ("SEPARATORS", "Are hyphens and spaces placed as the canonical form places them, and no more?",
     "nothing was retrieved (tier 0)", "2-bromobutane", "2-bromo-butane, iodo-ethane"),
]
for i, row in enumerate(card_rows):
    rr = 8 + i
    fill = FEATFILL if i < 4 else FORMFILL
    for j, v in enumerate(row, start=1):
        put(ws, rr, j, v, BODYB if j == 1 else BODY, wrap=True, fill=fill, border=True)
    ws.row_dimensions[rr].height = 30

put(ws, 15, 1, "Rows 1 to 4 are FEATURE KNOWLEDGE (blue). Rows 5 to 7 are PRODUCTION FORM (orange). "
               "They are scored and reported separately and are never pooled.", GREY, wrap=True)
ws.merge_cells("A15:E15")

put(ws, 17, 1, "The one decision rule", H2)
put(ws, 18, 1, "Credit a response that identifies the correct chemical entity but writes it non-canonically. "
               "Give no credit where the response denotes a different molecule.", BODYB, wrap=True)
ws.merge_cells("A18:E18")
put(ws, 19, 1, "Score each row independently against the target. A wrong parent does not automatically "
               "cost the class or substituent rows.", BODY, wrap=True)
ws.merge_cells("A19:E19")

put(ws, 21, 1, "The tier ladder", H2)
tier_rows = [
    (4, "Exact", "Matches the IUPAC written form. The current strict rule.", "1-chloroheptane"),
    (3, "Valid variant", "Chemically unambiguous and correct, in a non-preferred form.", "3-hexene, 1-iodoethane"),
    (2, "Complete but malformed", "Right entity, separator or spelling slip only.", "2-bromo-butane, 1-flourohexane, hept-2-one"),
    (1, "Partial retrieval", "Some rows correct, but the name is incomplete or denotes a different compound.", "fluoro-hexane, octyl methanoate"),
    (0, "No retrieval", "Blank, or nothing recoverable. Feature rows all 0, form rows all NA.", "chlorine, septane, (blank)"),
]
for j, h in enumerate(["Tier", "Name", "Definition", "Seen in the data"], start=1):
    put(ws, 22, j, h, H2, fill=HDRFILL, border=True)
for i, row in enumerate(tier_rows):
    rr = 23 + i
    for j, v in enumerate(row, start=1):
        put(ws, rr, j, v, BODY, wrap=True, border=True)
    ws.row_dimensions[rr].height = 26
put(ws, 28, 1, "Tier and row scores are allowed to disagree, and their disagreement is informative. "
               "octyl methanoate has every segment (feature score 1.00) and denotes a different ester (tier 1).",
    GREY, wrap=True)
ws.merge_cells("A28:E28")

put(ws, 30, 1, "Order of work", H2)
steps = [
    "1. Run the mechanical sweep first. Exact matches, separator-only near-misses and blanks are decided by the "
    "script, not by you, and need no coder. On the six completers that cleared 19 of 72 responses.",
    "2. Code the rest blind. The coding card carries a hashed coder key and no schedule label. Never code from a "
    "sheet that shows spaced or massed, because the rubric is applied to a within-subject contrast.",
    "3. Two coders, then adjudicate. Flag anything you would not defend in the Adjudicate column and say why in "
    "the Note column. On the six completers 6 of 72 needed adjudication.",
    "4. Report both rules side by side. Strict stays confirmatory. Segment credit enters as a pre-specified "
    "secondary and sensitivity layer, and only the spaced-minus-massed difference is comparable across rules.",
]
for i, t in enumerate(steps):
    put(ws, 31 + i, 1, t, BODY, wrap=True)
    ws.merge_cells(start_row=31 + i, start_column=1, end_row=31 + i, end_column=5)
    ws.row_dimensions[31 + i].height = 26

put(ws, 36, 1, "Tabs in this workbook", H2)
tabs = [
    ("Coding card", "Blank, fillable, coder-blind. One example row in grey. Type into the yellow cells only."),
    ("Coded 72", "All 72 whole-name responses from the six completers, coded by one unblinded coder. "
                 "A feasibility demonstration, not a result."),
    ("Summary", "Every table computed by formula from Coded 72. Nothing here is typed in."),
    ("Tags", "The controlled vocabulary. Use these strings only, separated by semicolons."),
]
for i, (a, b) in enumerate(tabs):
    put(ws, 37 + i, 1, a, BODYB, border=True)
    put(ws, 37 + i, 2, b, BODY, wrap=True, border=True)
    ws.merge_cells(start_row=37 + i, start_column=2, end_row=37 + i, end_column=5)
    ws.row_dimensions[37 + i].height = 26

for col, w in zip("ABCDE", [14, 46, 26, 30, 34]):
    ws.column_dimensions[col].width = w

# ════════════════════════════════════════════════════════════════════════════
# 2. TAGS
# ════════════════════════════════════════════════════════════════════════════
tg = wb.create_sheet("Tags")
tg.sheet_view.showGridLines = False
put(tg, 1, 1, "Controlled error-tag vocabulary", H1)
put(tg, 2, 1, "Multi-label and descriptive, not a score. Separate several tags with a semicolon. "
              "Do not invent tags: if none fits, flag the response for adjudication and say so in the Note.",
    GREY, wrap=True)
tg.merge_cells("A2:C2")
TAGS = [
    ("Feature failure", "wrong-root", "Parent chain length or ring identity is wrong."),
    ("Feature failure", "wrong-suffix-class", "Wrong functional class."),
    ("Feature failure", "wrong-substituent", "Wrong substituent prefix or branch descriptor."),
    ("Feature failure", "missing-substituent", "Substituent absent altogether."),
    ("Feature failure", "missing-locant", "A required locant is absent."),
    ("Feature failure", "wrong-locant-value", "Locant present but names the wrong position."),
    ("Feature failure", "missing-ring", "Chain length retrieved, the cyclo or aromatic marker lost."),
    ("Feature failure", "numeral-form-error", "Wrong numeral morpheme, for example sept for hept."),
    ("Form failure", "ordering", "Segments in the wrong sequence."),
    ("Form failure", "ester-inversion", "Acyl and alkyl parts of an ester exchanged."),
    ("Form failure", "elision", "The -an- linker is missing, for example hept-2-one."),
    ("Form failure", "word-boundary", "Word inserted or removed, for example methyloctanoate."),
    ("Form failure", "separator-count", "Hyphen inserted or removed."),
    ("Form failure", "orthographic", "Spelling slip in a correctly retrieved segment."),
    ("Variant", "legacy-locant-style", "Locant in the legacy position, for example 3-hexene."),
    ("Variant", "redundant-locant", "Locant present but unnecessary and harmless."),
    ("Variant", "alternative-nomenclature", "A different but valid naming system."),
    ("Non-response", "truncated", "Answer cut off mid-name."),
    ("Non-response", "blank", "Nothing typed."),
    ("Non-response", "timeout", "The 20-second timer expired."),
]
for j, h in enumerate(["Group", "Tag", "Meaning"], start=1):
    put(tg, 4, j, h, H2, fill=HDRFILL, border=True)
for i, row in enumerate(TAGS):
    for j, v in enumerate(row, start=1):
        put(tg, 5 + i, j, v, BODY, border=True)
for col, w in zip("ABC", [18, 26, 62]):
    tg.column_dimensions[col].width = w

# ════════════════════════════════════════════════════════════════════════════
# 3. CODING CARD (blank, coder-blind)
# ════════════════════════════════════════════════════════════════════════════
cc = wb.create_sheet("Coding card")
CC_HDR = ["Coder key", "Item id", "Target (correct name)", "Response (as typed)",
          "PARENT", "CLASS", "SUBST", "LOCANTS", "ORDER", "SPELLING", "SEPARATORS",
          "n feature", "Feature score", "n form", "Form score",
          "Tier", "Tags", "Adjudicate", "Note"]
put(cc, 1, 1, "Q-full coding card", H1)
put(cc, 2, 1, "Type into the yellow cells only. Grey row 4 is a worked example, do not code over it. "
              "Mark each of the seven rows 1, 0 or NA. The score columns fill themselves.", GREY, wrap=True)
cc.merge_cells("A2:S2")
for j, h in enumerate(CC_HDR, start=1):
    c = put(cc, 3, j, h, H2, fill=HDRFILL, border=True, wrap=True)
    if 5 <= j <= 8:
        c.fill = FEATFILL
    if 9 <= j <= 11:
        c.fill = FORMFILL
cc.row_dimensions[3].height = 30

example = ["a1b2c3d4", "C2-S2-oate-full", "methyl octanoate", "methyloctanoate",
           1, 1, 1, "NA", 0, 1, 0, None, None, None, None,
           2, "word-boundary", "N", "Ester written as one word. Every segment present."]
for j, v in enumerate(example, start=1):
    if j in (12, 13, 14, 15):
        continue
    put(cc, 4, j, v, BODY, fill=EXFILL, border=True)
cc["L4"] = "=COUNT(E4:H4)"
cc["M4"] = '=IFERROR(AVERAGE(E4:H4),"")'
cc["N4"] = "=COUNT(I4:K4)"
cc["O4"] = '=IFERROR(AVERAGE(I4:K4),"")'
for col in "LMNO":
    cc[f"{col}4"].font = BODY
    cc[f"{col}4"].fill = EXFILL
    cc[f"{col}4"].border = BOX

FIRST, LAST = 5, 84
for r in range(FIRST, LAST + 1):
    for j in range(1, 20):
        c = cc.cell(row=r, column=j)
        c.font = BODY
        c.border = BOX
        if j in (12, 13, 14, 15):
            continue
        c.fill = INFILL
    cc.cell(row=r, column=12, value=f"=COUNT(E{r}:H{r})")
    cc.cell(row=r, column=13, value=f'=IFERROR(AVERAGE(E{r}:H{r}),"")')
    cc.cell(row=r, column=14, value=f"=COUNT(I{r}:K{r})")
    cc.cell(row=r, column=15, value=f'=IFERROR(AVERAGE(I{r}:K{r}),"")')
    for col in "MO":
        cc[f"{col}{r}"].number_format = "0.00"

dv_slot = DataValidation(type="list", formula1='"1,0,NA"', allow_blank=True)
dv_tier = DataValidation(type="list", formula1='"0,1,2,3,4"', allow_blank=True)
dv_adj = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
cc.add_data_validation(dv_slot)
cc.add_data_validation(dv_tier)
cc.add_data_validation(dv_adj)
for col in "EFGHIJK":
    dv_slot.add(f"{col}{FIRST}:{col}{LAST}")
dv_tier.add(f"P{FIRST}:P{LAST}")
dv_adj.add(f"R{FIRST}:R{LAST}")

cc["E3"].comment = Comment(
    "Mark NA only where the target does not call for the row.\n"
    "SUBST: NA when the target carries no substituent.\n"
    "LOCANTS: NA when the target requires no locant.\n"
    "ORDER, SPELLING, SEPARATORS: NA only on a tier-0 response, where nothing was retrieved.",
    "rubric v1")
widths = [12, 20, 24, 24, 10, 10, 10, 11, 10, 11, 13, 10, 12, 9, 10, 7, 30, 11, 44]
for j, w in enumerate(widths, start=1):
    cc.column_dimensions[get_column_letter(j)].width = w
cc.freeze_panes = "E4"

# ════════════════════════════════════════════════════════════════════════════
# 4. CODED 72
# ════════════════════════════════════════════════════════════════════════════
cd = wb.create_sheet("Coded 72")
CD_HDR = ["pid", "Schedule", "List", "Item id", "Target", "Response",
          "PARENT", "CLASS", "SUBST", "LOCANTS", "ORDER", "SPELLING", "SEPARATORS",
          "n feature", "Feature score", "n form", "Form score",
          "Tier", "Exact", "Tags", "Timeout", "Adjudicate", "Note"]
put(cd, 1, 1, "All 72 whole-name responses, six completers, coded under the simplified card", H1)
put(cd, 2, 1, "One coder, not blind to condition, not pre-specified. This is a feasibility demonstration "
              "of the instrument, not a result, and no number here belongs in Chapter 4.", GREY, wrap=True)
cd.merge_cells("A2:W2")
for j, h in enumerate(CD_HDR, start=1):
    c = put(cd, 3, j, h, H2, fill=HDRFILL, border=True, wrap=True)
    if 7 <= j <= 10:
        c.fill = FEATFILL
    if 11 <= j <= 13:
        c.fill = FORMFILL
cd.row_dimensions[3].height = 30

for i, (_, r) in enumerate(coded.iterrows()):
    rr = 4 + i
    vals = [r.pid, r.schedule, r["list"], r.item_id, r.target, r.response,
            r.PARENT, r.CLASS, r.SUBST, r.LOCANTS, r.ORDER, r.SPELLING, r.SEPARATORS,
            None, None, None, None, r.tier, None, r.tags, r.timeout, r.adjudicate, r.note]
    for j, v in enumerate(vals, start=1):
        c = put(cd, rr, j, v, BODY, border=True)
    cd.cell(row=rr, column=14, value=f"=COUNT(G{rr}:J{rr})").border = BOX
    cd.cell(row=rr, column=15, value=f'=IFERROR(AVERAGE(G{rr}:J{rr}),"")').border = BOX
    cd.cell(row=rr, column=16, value=f"=COUNT(K{rr}:M{rr})").border = BOX
    cd.cell(row=rr, column=17, value=f'=IFERROR(AVERAGE(K{rr}:M{rr}),"")').border = BOX
    cd.cell(row=rr, column=19, value=f"=IF(R{rr}=4,1,0)").border = BOX
    for col in ("N", "O", "P", "Q", "S"):
        cd[f"{col}{rr}"].font = BODY
    for col in ("O", "Q"):
        cd[f"{col}{rr}"].number_format = "0.00"
LAST_CD = 3 + len(coded)

widths = [11, 11, 7, 20, 22, 24, 10, 10, 10, 11, 10, 11, 13, 10, 12, 9, 10, 7, 8, 30, 10, 11, 46]
for j, w in enumerate(widths, start=1):
    cd.column_dimensions[get_column_letter(j)].width = w
cd.freeze_panes = "G4"

# ════════════════════════════════════════════════════════════════════════════
# 5. SUMMARY  (all formulas, nothing typed)
# ════════════════════════════════════════════════════════════════════════════
sm = wb.create_sheet("Summary")
sm.sheet_view.showGridLines = False
D = f"'Coded 72'"
R1, R2 = 4, LAST_CD
put(sm, 1, 1, "Summary", H1)
put(sm, 2, 1, "Every figure on this sheet is computed by formula from the Coded 72 tab. Nothing here is typed in. "
              "Six participants, six whole-name items per condition, so no contrast is distinguishable from zero.",
    GREY, wrap=True)
sm.merge_cells("A2:O2")

pids = sorted(coded.pid.unique())

# Block A ─ participant x condition
put(sm, 4, 1, "A. Participant by condition", H2)
heads = ["Participant", "Exact, spaced", "Exact, massed", "Feature, spaced", "Feature, massed",
         "Form, spaced", "Form, massed", "Timeouts, spaced", "Timeouts, massed",
         "Diff, exact", "Diff, feature", "Diff, form",
         "Feature ex-timeout, spaced", "Feature ex-timeout, massed", "Diff, ex-timeout"]
for j, h in enumerate(heads, start=1):
    put(sm, 5, j, h, H2, fill=HDRFILL, border=True, wrap=True)
sm.row_dimensions[5].height = 28
srccol = {"Exact": "S", "Feature": "O", "Form": "Q"}
for i, p in enumerate(pids):
    rr = 6 + i
    put(sm, rr, 1, p, BODY, border=True)
    j = 2
    for meas in ("Exact", "Feature", "Form"):
        for sch in ("spaced", "massed"):
            f = (f'=AVERAGEIFS({D}!${srccol[meas]}${R1}:${srccol[meas]}${R2},'
                 f'{D}!$A${R1}:$A${R2},$A{rr},{D}!$B${R1}:$B${R2},"{sch}")')
            c = sm.cell(row=rr, column=j, value=f)
            c.font = BODY
            c.border = BOX
            c.number_format = "0.0%"
            j += 1
    for sch in ("spaced", "massed"):
        f = (f'=SUMIFS({D}!$U${R1}:$U${R2},{D}!$A${R1}:$A${R2},$A{rr},'
             f'{D}!$B${R1}:$B${R2},"{sch}")')
        c = sm.cell(row=rr, column=j, value=f)
        c.font = BODY
        c.border = BOX
        c.number_format = "0"
        j += 1
    for j, (cs, cm) in ((10, ("B", "C")), (11, ("D", "E")), (12, ("F", "G"))):
        c = sm.cell(row=rr, column=j, value=f"={cs}{rr}-{cm}{rr}")
        c.font = BODY
        c.border = BOX
        c.number_format = "0.000"
    for j, sch in ((13, "spaced"), (14, "massed")):
        f = (f'=AVERAGEIFS({D}!$O${R1}:$O${R2},{D}!$A${R1}:$A${R2},$A{rr},'
             f'{D}!$B${R1}:$B${R2},"{sch}",{D}!$U${R1}:$U${R2},0)')
        c = sm.cell(row=rr, column=j, value=f)
        c.font = BODY
        c.border = BOX
        c.number_format = "0.0%"
    c = sm.cell(row=rr, column=15, value=f"=M{rr}-N{rr}")
    c.font = BODY
    c.border = BOX
    c.number_format = "0.000"
AR1, AR2 = 6, 6 + len(pids) - 1

# Block B ─ contrasts
put(sm, 13, 1, "B. Participant-level contrasts", H2)
for j, h in enumerate(["Measure", "Spaced", "Massed", "Difference", "SD of difference",
                       "Paired t, two-tailed p", "Favouring spaced"], start=1):
    put(sm, 14, j, h, H2, fill=HDRFILL, border=True, wrap=True)
sm.row_dimensions[14].height = 28
contrasts = [("Exact match (current strict rule)", "B", "C", "J"),
             ("Feature score", "D", "E", "K"),
             ("Form score", "F", "G", "L")]
for i, (label, cs, cm, cdif) in enumerate(contrasts):
    rr = 15 + i
    put(sm, rr, 1, label, BODY, border=True)
    for j, col in ((2, cs), (3, cm)):
        c = sm.cell(row=rr, column=j, value=f"=AVERAGE({col}{AR1}:{col}{AR2})")
        c.font = BODY
        c.border = BOX
        c.number_format = "0.0%"
    c = sm.cell(row=rr, column=4, value=f"=B{rr}-C{rr}")
    c.font = BODYB
    c.border = BOX
    c.number_format = "0.0%;(0.0%);-"
    c = sm.cell(row=rr, column=5, value=f"=STDEV({cdif}{AR1}:{cdif}{AR2})")
    c.font = BODY
    c.border = BOX
    c.number_format = "0.000"
    c = sm.cell(row=rr, column=6,
                value=f"=TTEST({cs}{AR1}:{cs}{AR2},{cm}{AR1}:{cm}{AR2},2,1)")
    c.font = BODY
    c.border = BOX
    c.number_format = "0.000"
    c = sm.cell(row=rr, column=7,
                value=f"=SUMPRODUCT(--({cs}{AR1}:{cs}{AR2}>{cm}{AR1}:{cm}{AR2}))"
                      f'&" of "&COUNT({cs}{AR1}:{cs}{AR2})')
    c.font = BODY
    c.border = BOX

rr = 18
put(sm, rr, 1, "Feature score, timed-out trials excluded", BODY, border=True)
for j, col in ((2, "M"), (3, "N")):
    c = sm.cell(row=rr, column=j, value=f"=AVERAGE({col}{AR1}:{col}{AR2})")
    c.font = BODY
    c.border = BOX
    c.number_format = "0.0%"
c = sm.cell(row=rr, column=4, value=f"=B{rr}-C{rr}")
c.font = BODYB
c.border = BOX
c.number_format = "0.0%;(0.0%);-"
c = sm.cell(row=rr, column=5, value=f"=STDEV(O{AR1}:O{AR2})")
c.font = BODY
c.border = BOX
c.number_format = "0.000"
c = sm.cell(row=rr, column=6, value=f"=TTEST(M{AR1}:M{AR2},N{AR1}:N{AR2},2,1)")
c.font = BODY
c.border = BOX
c.number_format = "0.000"
c = sm.cell(row=rr, column=7,
            value=f"=SUMPRODUCT(--(M{AR1}:M{AR2}>N{AR1}:N{AR2}))"
                  f'&" of "&COUNT(M{AR1}:M{AR2})')
c.font = BODY
c.border = BOX
put(sm, 19, 1, "The sign flip between the feature score and the same score with timeouts excluded is produced "
               "by seven trials. See the timeout block below.", GREY, wrap=True)
sm.merge_cells("A19:G19")

put(sm, 20, 1, "Timed-out trials: whole-name format", H2)
for j, h in enumerate(["", "Spaced", "Massed"], start=1):
    put(sm, 21, j, h, H2, fill=HDRFILL, border=True)
put(sm, 22, 1, "Trials that ran the 20-second timer out", BODY, border=True)
for j, sch in ((2, "spaced"), (3, "massed")):
    c = sm.cell(row=22, column=j,
                value=f'=SUMIFS({D}!$U${R1}:$U${R2},{D}!$B${R1}:$B${R2},"{sch}")'
                      f'&" of "&COUNTIFS({D}!$B${R1}:$B${R2},"{sch}")')
    c.font = BODYB
    c.border = BOX
put(sm, 23, 1, "Participants showing the asymmetry", BODY, border=True)
for j, (cs, cm) in ((2, ("H", "I")), (3, ("I", "H"))):
    c = sm.cell(row=23, column=j,
                value=f"=SUMPRODUCT(--({cs}{AR1}:{cs}{AR2}>{cm}{AR1}:{cm}{AR2}))"
                      f'&" of "&COUNT({cs}{AR1}:{cs}{AR2})')
    c.font = BODY
    c.border = BOX

# Block C ─ tier distribution
put(sm, 25, 1, "C. Tier distribution", H2)
for j, h in enumerate(["Tier", "Spaced", "Massed", "Total"], start=1):
    put(sm, 26, j, h, H2, fill=HDRFILL, border=True)
tiernames = {4: "4  Exact", 3: "3  Valid variant", 2: "2  Complete but malformed",
             1: "1  Partial retrieval", 0: "0  No retrieval"}
for i, t in enumerate([4, 3, 2, 1, 0]):
    rr = 27 + i
    put(sm, rr, 1, tiernames[t], BODY, border=True)
    for j, sch in ((2, "spaced"), (3, "massed")):
        c = sm.cell(row=rr, column=j,
                    value=f'=COUNTIFS({D}!$R${R1}:$R${R2},{t},{D}!$B${R1}:$B${R2},"{sch}")')
        c.font = BODY
        c.border = BOX
    c = sm.cell(row=rr, column=4, value=f"=B{rr}+C{rr}")
    c.font = BODY
    c.border = BOX
put(sm, 32, 1, "Total", BODYB, border=True)
for col in "BCD":
    c = sm.cell(row=32, column="ABCD".index(col) + 1, value=f"=SUM({col}27:{col}31)")
    c.font = BODYB
    c.border = BOX

# Block D ─ per item
put(sm, 34, 1, "D. Per whole-name item", H2)
for j, h in enumerate(["Target", "List", "Exact-match rate", "Mean feature score",
                       "Mean form score", "Responses"], start=1):
    put(sm, 35, j, h, H2, fill=HDRFILL, border=True, wrap=True)
items = coded[["target", "list", "item_id"]].drop_duplicates().sort_values(["list", "target"])
for i, (_, it) in enumerate(items.iterrows()):
    rr = 36 + i
    put(sm, rr, 1, it.target, BODY, border=True)
    put(sm, rr, 2, it["list"], BODY, border=True)
    for j, col in ((3, "S"), (4, "O"), (5, "Q")):
        c = sm.cell(row=rr, column=j,
                    value=f'=AVERAGEIFS({D}!${col}${R1}:${col}${R2},'
                          f'{D}!$E${R1}:$E${R2},$A{rr})')
        c.font = BODY
        c.border = BOX
        c.number_format = "0.0%"
    c = sm.cell(row=rr, column=6,
                value=f'=COUNTIFS({D}!$E${R1}:$E${R2},$A{rr})')
    c.font = BODY
    c.border = BOX

put(sm, 49, 1, "E. Coding workload", H2)
work = [
    ("Responses coded", f'=COUNT({D}!$R${R1}:$R${R2})'),
    ("Exact matches, decided mechanically", f'=COUNTIFS({D}!$R${R1}:$R${R2},4)'),
    ("Blanks, decided mechanically", f'=COUNTIFS({D}!$R${R1}:$R${R2},0,{D}!$F${R1}:$F${R2},"(blank)")'),
    ("Flagged for adjudication", f'=COUNTIFS({D}!$V${R1}:$V${R2},"Y")'),
    ("Scored wrong under the strict rule", f'=COUNTIFS({D}!$S${R1}:$S${R2},0)'),
    ("  of those, with at least one feature row correct",
     f'=SUMPRODUCT(--({D}!$S${R1}:$S${R2}=0),--({D}!$O${R1}:$O${R2}>0))'),
    ("  of those, with every feature row correct",
     f'=SUMPRODUCT(--({D}!$S${R1}:$S${R2}=0),--({D}!$O${R1}:$O${R2}=1))'),
]
for i, (label, f) in enumerate(work):
    rr = 50 + i
    put(sm, rr, 1, label, BODY, border=True)
    c = sm.cell(row=rr, column=2, value=f)
    c.font = BODYB
    c.border = BOX

for col, w in zip("ABCDEFGHIJKLMNO", [40, 16, 16, 17, 17, 22, 18, 18, 18, 14, 14, 14, 17, 17, 15]):
    sm.column_dimensions[col].width = w

wb.save("Qfull_coding_card_v1.xlsx")
print("written")
